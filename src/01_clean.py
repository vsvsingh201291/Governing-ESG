#!/usr/bin/env python3
"""
Governing ESG - Step 1: clean and assemble the analytical datasets.

Reads the legacy thesis source files (Excel) and produces a single tidy
JSON file consumed by the web app. Dependency-light: standard library +
openpyxl only, so the whole pipeline is reproducible with `pip install openpyxl`.

Outputs: outputs/app_data.json
"""
import json
import math
import os
import statistics as st
from datetime import datetime

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
DATA = os.path.join(ROOT, "data")
OUT = os.path.join(ROOT, "outputs")
os.makedirs(OUT, exist_ok=True)


def num(x):
    return x if isinstance(x, (int, float)) and not isinstance(x, bool) else None


def to_num(x):
    """Coerce numeric-looking strings (some Excel cells arrive as text)."""
    if isinstance(x, (int, float)) and not isinstance(x, bool):
        return float(x)
    if isinstance(x, str):
        try:
            return float(x.strip())
        except ValueError:
            return None
    return None


def rnd(x, d=4):
    return None if x is None else round(float(x), d)


# --------------------------------------------------------------------------
# 1. Firm panel (h2) - 206 companies x 6 years (2016-2021)
#    cols: A Dates | B board size | C age range | D women directors
#          E women executives | F pct women board | G esg-linked comp
#          H esg-linked bonus | I board tenure | J ESG disclosure score
# --------------------------------------------------------------------------
def load_firm_panel():
    wb = openpyxl.load_workbook(os.path.join(DATA, "pay and gender diversity data.xlsx"),
                                read_only=True, data_only=True)
    rows = list(wb["h2"].iter_rows(values_only=True))[1:]
    wb.close()

    # Rows run year-ascending within each company; a new company starts
    # when the year resets (e.g. 2016 follows 2021).
    blocks, cur, prev = [], [], None
    for r in rows:
        if r[0] is None:
            continue
        if prev is not None and r[0] < prev and cur:
            blocks.append(cur)
            cur = []
        cur.append(r)
        prev = r[0]
    if cur:
        blocks.append(cur)

    panel = []
    for i, b in enumerate(blocks, start=1):
        cid = "C%03d" % i
        for r in b:
            if num(r[0]) is None:
                continue
            panel.append({
                "company_id": cid,
                "company_no": i,
                "year": int(r[0]),
                "board_size": num(r[1]),
                "age_range": num(r[2]),
                "women_dir": num(r[3]),
                "women_exec": num(r[4]),
                "pct_women": None if num(r[5]) is None else rnd(r[5], 3),
                "esg_comp": num(r[6]),
                "esg_bonus": num(r[7]),
                "tenure": num(r[8]),
                "esg_disclosure": None if num(r[9]) is None else rnd(r[9], 3),
            })
    return panel


# --------------------------------------------------------------------------
# 2. H1 panel - pooled Year | ESG score | compensation
# --------------------------------------------------------------------------
def load_h1_panel():
    wb = openpyxl.load_workbook(os.path.join(DATA, "pay and gender diversity data.xlsx"),
                                read_only=True, data_only=True)
    rows = list(wb["H1"].iter_rows(values_only=True))[1:]
    wb.close()
    out = []
    for r in rows:
        if num(r[0]) is None or num(r[1]) is None:
            continue
        out.append({"year": int(r[0]), "esg_score": rnd(r[1], 3),
                    "compensation": None if num(r[2]) is None else rnd(r[2], 2)})
    return out


# --------------------------------------------------------------------------
# 3. Country panel - GDP growth + 5 World Bank governance estimates
# --------------------------------------------------------------------------
def load_country_panel():
    wb = openpyxl.load_workbook(os.path.join(DATA, "gdp data.xlsx"),
                                read_only=True, data_only=True)
    rows = list(wb["Sheet3"].iter_rows(values_only=True))[1:]
    wb.close()
    out = []
    for r in rows:
        if num(r[0]) is None:
            continue
        out.append({
            "gdp_growth": rnd(r[0], 3),
            "control_corruption": rnd(r[1]),
            "govt_effectiveness": rnd(r[2]),
            "political_stability": rnd(r[3]),
            "regulatory_quality": rnd(r[4]),
            "rule_of_law": rnd(r[5]),
        })
    return out


# --------------------------------------------------------------------------
# 4. UK named companies - FTSE board roster + remuneration/dividend series
#    Both FTSE files share the same 205-company / director-row layout.
# --------------------------------------------------------------------------
def parse_ftse_blocks(path, sheet="Results"):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(wb[sheet].iter_rows(values_only=True))
    wb.close()
    header, data = rows[0], rows[1:]
    blocks, cur = [], None
    for r in data:
        if r[1]:
            cur = {"name": str(r[1]).strip(), "rows": []}
            blocks.append(cur)
        if cur is not None:
            cur["rows"].append(r)
    return header, blocks


def load_uk_companies():
    hdr_bd, bd_blocks = parse_ftse_blocks(os.path.join(DATA, "board diversity data ftse.xlsx"))
    hdr_rem, rem_blocks = parse_ftse_blocks(os.path.join(DATA, "ftse remuneration.xlsx"))

    # remuneration file: col index (0-based) -> label. Find year columns dynamically.
    rem_years = {}
    for idx, h in enumerate(hdr_rem):
        if not h:
            continue
        label = str(h).replace("\n", " ")
        for y in range(2003, 2022):
            if label.startswith("Directors' Remuneration") and str(y) in label:
                rem_years[y] = idx
        for y in range(2003, 2022):
            if label.startswith("Dividends") and str(y) in label:
                rem_years.setdefault("div_%d" % y, idx)

    rem_by_name = {b["name"]: b for b in rem_blocks}
    out = []
    for b in bd_blocks:
        rows = b["rows"]
        genders = [r[18] for r in rows if r[18] in ("M", "F")]
        ages = [to_num(r[20]) for r in rows if to_num(r[20])]
        n_f = sum(1 for g in genders if g == "F")
        pct_f = (100.0 * n_f / len(genders)) if genders else None

        rem = rem_by_name.get(b["name"])
        rem_series = {}
        div_series = {}
        if rem and rem["rows"]:
            first = rem["rows"][0]
            for y, idx in rem_years.items():
                if isinstance(y, int):
                    v = to_num(first[idx])
                    if v is not None:
                        rem_series[str(y)] = rnd(v, 0)
                else:
                    v = to_num(first[idx])
                    if v is not None:
                        div_series[str(y).replace("div_", "")] = rnd(v, 0)

        rem_years_sorted = sorted(rem_series)
        rem_latest = rem_series[rem_years_sorted[-1]] if rem_years_sorted else None

        out.append({
            "name": b["name"],
            "records": len(rows),
            "n_male": len(genders) - n_f,
            "n_female": n_f,
            "pct_female": None if pct_f is None else rnd(pct_f, 2),
            "avg_age": None if not ages else rnd(st.mean(ages), 1),
            "min_age": None if not ages else rnd(min(ages), 0),
            "max_age": None if not ages else rnd(max(ages), 0),
            "age_range": None if not ages else rnd(max(ages) - min(ages), 0),
            "rem_latest": rem_latest,
            "rem_latest_year": rem_years_sorted[-1] if rem_years_sorted else None,
            "remuneration": rem_series,
            "dividends": div_series,
        })

    # research-derived signals: board gender diversity vs the peer set
    valid = sorted(c["pct_female"] for c in out if c["pct_female"] is not None)
    median_pf = st.median(valid) if valid else None
    for c in out:
        pf = c["pct_female"]
        if pf is None:
            c["diversity_score"] = None
            c["uplift"] = None
            continue
        # percentile rank (0-100) of % female directors within the 205-company peer set
        rank = sum(1 for v in valid if v <= pf)
        c["diversity_score"] = int(round(100.0 * rank / len(valid)))
        # implied disclosure uplift vs a board with 0% female directors,
        # using the H2 coefficient (+0.1851 points per percentage point)
        c["uplift"] = rnd(0.1851 * pf, 1)
    summary_peer = {"median_pct_female": rnd(median_pf, 1) if median_pf is not None else None,
                    "with_data": len(valid)}
    return out, summary_peer


# --------------------------------------------------------------------------
# 5. OLS (normal equations, pure python) for the H2 disclosure model
# --------------------------------------------------------------------------
def solve(a, b):
    """Gaussian elimination with partial pivoting: solve a x = b."""
    n = len(a)
    m = [row[:] + [b[i]] for i, row in enumerate(a)]
    for col in range(n):
        piv = max(range(col, n), key=lambda r: abs(m[r][col]))
        m[col], m[piv] = m[piv], m[col]
        if abs(m[col][col]) < 1e-12:
            continue
        pv = m[col][col]
        m[col] = [v / pv for v in m[col]]
        for r in range(n):
            if r != col and m[r][col]:
                f = m[r][col]
                m[r] = [v - f * m[col][k] for k, v in enumerate(m[r])]
    return [m[i][n] for i in range(n)]


def invert(a):
    n = len(a)
    inv = []
    for i in range(n):
        e = [1.0 if j == i else 0.0 for j in range(n)]
        inv.append(solve([row[:] for row in a], e))
    # inv rows are solutions; transpose to get inverse
    return [[inv[j][i] for j in range(n)] for i in range(n)]


def norm_cdf(z):
    return 0.5 * (1.0 + math.erf(z / math.sqrt(2.0)))


def ols(X, y, names):
    n, k = len(X), len(X[0])
    XtX = [[sum(X[r][i] * X[r][j] for r in range(n)) for j in range(k)] for i in range(k)]
    Xty = [sum(X[r][i] * y[r] for r in range(n)) for i in range(k)]
    beta = solve([row[:] for row in XtX], Xty)
    yhat = [sum(X[r][j] * beta[j] for j in range(k)) for r in range(n)]
    resid = [y[r] - yhat[r] for r in range(n)]
    sse = sum(e * e for e in resid)
    ybar = sum(y) / n
    sst = sum((v - ybar) ** 2 for v in y)
    r2 = 1 - sse / sst if sst else None
    sigma2 = sse / (n - k)
    XtX_inv = invert([row[:] for row in XtX])
    params = []
    for j in range(k):
        se = math.sqrt(max(sigma2 * XtX_inv[j][j], 0))
        t = beta[j] / se if se else 0
        p = 2 * (1 - norm_cdf(abs(t)))
        sig = "***" if p < 0.01 else "**" if p < 0.05 else "*" if p < 0.1 else ""
        params.append({"term": names[j], "coef": rnd(beta[j], 4), "se": rnd(se, 4),
                       "t": rnd(t, 3), "p": rnd(p, 4), "sig": sig})
    return {"n": n, "r2": rnd(r2, 4), "sigma": rnd(math.sqrt(sigma2), 4),
            "names": names,
            "xtx_inv": [[rnd(v, 8) for v in row] for row in XtX_inv],
            "params": params,
            "yhat": yhat, "resid": resid}


def build_model(panel):
    rows = [r for r in panel if r["esg_disclosure"] is not None]
    names = ["const", "board_size", "age_range", "women_dir", "women_exec",
             "pct_women", "esg_comp", "esg_bonus", "tenure"]
    X, y = [], []
    for r in rows:
        if any(r[c] is None for c in ["board_size", "age_range", "women_dir",
                                      "women_exec", "pct_women", "esg_comp",
                                      "esg_bonus", "tenure"]):
            continue
        X.append([1.0, r["board_size"], r["age_range"], r["women_dir"], r["women_exec"],
                  r["pct_women"], r["esg_comp"], r["esg_bonus"], r["tenure"]])
        y.append(r["esg_disclosure"])
    model = ols(X, y, names)
    model.pop("yhat"); model.pop("resid")

    # Governance-driven score: coefficient-weighted drivers, min-max scaled 0-100
    drv = ["pct_women", "women_exec", "esg_bonus", "tenure", "age_range"]
    coef = {p["term"]: p["coef"] for p in model["params"]}
    scored = []
    allrows = [r for r in panel if r["esg_disclosure"] is not None]
    for r in allrows:
        s = sum(coef[t] * r[t] for t in drv if r[t] is not None)
        scored.append((r, s))
    lo = min(s for _, s in scored)
    hi = max(s for _, s in scored)
    for r, s in scored:
        r["gov_score"] = rnd(100 * (s - lo) / (hi - lo), 1) if hi > lo else 50.0
    return model


# --------------------------------------------------------------------------
def main():
    firm = load_firm_panel()
    h1 = load_h1_panel()
    country = load_country_panel()
    uk, peer = load_uk_companies()
    model = build_model(firm)

    # summary stats
    disclosures = [r["esg_disclosure"] for r in firm if r["esg_disclosure"] is not None]
    pw = [r["pct_women"] for r in firm if r["pct_women"] is not None]
    companies = sorted({r["company_id"] for r in firm})

    def mmean(key):
        vals = [r[key] for r in firm if r[key] is not None]
        return rnd(st.mean(vals), 2) if vals else None

    def share(key):
        return rnd(100.0 * sum(1 for r in firm if r[key] == 1) / len(firm), 1)

    # per-company averages
    by_cid = {}
    for r in firm:
        by_cid.setdefault(r["company_id"], []).append(r)
    company_avgs = []
    for cid in companies:
        rs = by_cid[cid]
        def avg(key):
            vals = [x[key] for x in rs if x[key] is not None]
            return rnd(st.mean(vals), 2) if vals else None
        company_avgs.append({
            "company_id": cid,
            "board_size": avg("board_size"), "pct_women": avg("pct_women"),
            "women_exec": avg("women_exec"), "tenure": avg("tenure"),
            "esg_disclosure": avg("esg_disclosure"), "gov_score": avg("gov_score"),
        })
    company_avgs.sort(key=lambda x: (x["gov_score"] is None, -(x["gov_score"] or 0)))

    corr = None
    xs = [r["pct_women"] for r in firm if r["pct_women"] is not None and r["esg_disclosure"] is not None]
    ys = [r["esg_disclosure"] for r in firm if r["pct_women"] is not None and r["esg_disclosure"] is not None]
    if xs:
        mx, my = st.mean(xs), st.mean(ys)
        num_ = sum((a - mx) * (b - my) for a, b in zip(xs, ys))
        den = (sum((a - mx) ** 2 for a in xs) ** 0.5) * (sum((b - my) ** 2 for b in ys) ** 0.5)
        corr = rnd(num_ / den, 3) if den else None

    payload = {
        "meta": {
            "title": "Governing ESG",
            "subtitle": "Does board governance drive ESG disclosure?",
            "generated": datetime.now().strftime("%Y-%m-%d"),
            "source": "MSc thesis dataset (Cranfield, 2022)",
            "firm_observations": len(firm),
            "firms": len(companies),
            "years": sorted({r["year"] for r in firm}),
            "uk_companies": len(uk),
        },
        "summary": {
            "esg_disclosure_mean": rnd(st.mean(disclosures), 2),
            "esg_disclosure_min": rnd(min(disclosures), 2),
            "esg_disclosure_max": rnd(max(disclosures), 2),
            "pct_women_mean": rnd(st.mean(pw), 2),
            "pct_women_vs_disclosure_corr": corr,
            "peer_median_pct_female": peer["median_pct_female"],
            "peers_with_data": peer["with_data"],
            "board_size_mean": mmean("board_size"),
            "age_range_mean": mmean("age_range"),
            "women_dir_mean": mmean("women_dir"),
            "women_exec_mean": mmean("women_exec"),
            "tenure_mean": mmean("tenure"),
            "esg_bonus_share": share("esg_bonus"),
            "esg_comp_share": share("esg_comp"),
        },
        "firm_panel": firm,
        "company_avgs": company_avgs,
        "h1_panel": h1,
        "country_panel": country,
        "uk_companies": uk,
        "model": model,
    }

    with open(os.path.join(OUT, "app_data.json"), "w") as f:
        json.dump(payload, f, separators=(",", ":"))

    size = os.path.getsize(os.path.join(OUT, "app_data.json"))
    print("firm observations:", len(firm), "| firms:", len(companies))
    print("h1:", len(h1), "| country:", len(country), "| uk:", len(uk))
    print("H2 model n=%s R2=%s" % (model["n"], model["r2"]))
    for p in model["params"]:
        print("   %-12s coef=%-9s se=%-8s t=%-7s p=%-8s %s" %
              (p["term"], p["coef"], p["se"], p["t"], p["p"], p["sig"]))
    print("pct_women vs disclosure corr:", corr)
    print("wrote outputs/app_data.json (%.1f KB)" % (size / 1024))


if __name__ == "__main__":
    main()
